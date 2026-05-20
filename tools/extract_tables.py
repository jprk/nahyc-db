#!/usr/bin/env python3
"""
Extracts tables from a multi-sheet XLSX file into individual CSV files.

Table structure assumptions:
- Each table starts with a header row where the first header cell = "Typ dokumentu"
  and the last header cell = "Poznámka"
- Tables can start at any column (not necessarily column 1)
- Multiple tables can sit side by side on the same row
- The table name (category) is in a merged cell 2 rows above the header
- A table ends at the first completely empty row (within the table's column range)
- Sheets may contain zero or more tables

Output: CSV files named "{sheet_name}_{table_name}.csv"
"""

import sys
import os
import re
import csv
from openpyxl import load_workbook


def sanitize_filename(name):
    """Remove or replace characters not suitable for filenames."""
    name = name.strip()
    name = re.sub(r'[\\/*?:"<>|]', '_', name)
    name = re.sub(r'\s+', ' ', name)
    return name


def find_table_name(ws, header_row, start_col, end_col):
    """
    Find the table name 2 rows above the header.
    Look for a merged cell or any non-empty cell in the table's column range.
    """
    name_row = header_row - 2
    if name_row < 1:
        return None

    # First check merged cells that overlap with name_row and our column range
    for merged_range in ws.merged_cells.ranges:
        if merged_range.min_row <= name_row <= merged_range.max_row:
            if merged_range.min_col <= end_col and merged_range.max_col >= start_col:
                val = ws.cell(row=merged_range.min_row, column=merged_range.min_col).value
                if val and str(val).strip():
                    return str(val).strip()

    # Fallback: scan columns in that row within the table's range
    for col in range(start_col, end_col + 1):
        val = ws.cell(row=name_row, column=col).value
        if val and str(val).strip():
            return str(val).strip()

    return None


def is_row_empty(ws, row, min_col, max_col):
    """Check if all cells in a row range are empty."""
    for col in range(min_col, max_col + 1):
        val = ws.cell(row=row, column=col).value
        if val is not None and str(val).strip() != '':
            return False
    return True


def find_all_tables(ws):
    """
    Scan the sheet for all table anchors.
    A table anchor is a cell with value "Typ dokumentu"; the table extends
    rightward until a cell with value "Poznámka" is found.
    Returns list of dicts: {header_row, start_col, end_col}
    """
    tables = []
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            val = ws.cell(row=row, column=col).value
            if val and str(val).strip() == "Typ dokumentu":
                # Find "Poznámka" to the right
                end_col = None
                for c in range(col + 1, max_col + 1):
                    v = ws.cell(row=row, column=c).value
                    if v and str(v).strip() == "Poznámka":
                        end_col = c
                        break
                if end_col:
                    tables.append({
                        'header_row': row,
                        'start_col': col,
                        'end_col': end_col,
                    })
    return tables


def extract_tables(input_file, output_dir=None):
    if output_dir is None:
        output_dir = os.path.splitext(input_file)[0] + '_csv'
    os.makedirs(output_dir, exist_ok=True)

    wb = load_workbook(input_file, data_only=True)
    total_tables = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n--- List: {sheet_name} ---")

        tables = find_all_tables(ws)
        if not tables:
            print("  Žádné tabulky nenalezeny.")
            continue

        for tbl in tables:
            hr = tbl['header_row']
            sc = tbl['start_col']
            ec = tbl['end_col']

            # Read header
            header = []
            for col in range(sc, ec + 1):
                val = ws.cell(row=hr, column=col).value
                header.append(str(val).strip() if val else '')

            # Get table name
            table_name = find_table_name(ws, hr, sc, ec)
            if not table_name:
                table_name = f"tabulka_r{hr}_c{sc}"

            print(f"  Tabulka: '{table_name}' (řádek {hr}, sloupce {sc}-{ec})")

            # Collect data rows until empty row
            data_rows = []
            for row in range(hr + 1, (ws.max_row or hr) + 1):
                if is_row_empty(ws, row, sc, ec):
                    break
                row_data = []
                for col in range(sc, ec + 1):
                    val = ws.cell(row=row, column=col).value
                    row_data.append('' if val is None else val)
                data_rows.append(row_data)

            # Write CSV with unique filename
            safe_sheet = sanitize_filename(sheet_name)
            safe_table = sanitize_filename(table_name)
            csv_filename = f"{safe_sheet}_{safe_table}.csv"
            csv_path = os.path.join(output_dir, csv_filename)
            counter = 1
            while os.path.exists(csv_path):
                csv_filename = f"{safe_sheet}_{safe_table}{counter}.csv"
                csv_path = os.path.join(output_dir, csv_filename)
                counter += 1

            with open(csv_path, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f, delimiter=';')
                writer.writerow(header)
                writer.writerows(data_rows)

            print(f"    -> {csv_filename} ({len(data_rows)} řádků)")
            total_tables += 1

    wb.close()
    print(f"\nHotovo. Celkem exportováno {total_tables} tabulek do: {output_dir}")


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(f"Použití: python {sys.argv[0]} <vstupní_soubor.xlsx> [výstupní_adresář]")
        sys.exit(1)

    input_file = sys.argv[1]
    output_dir = sys.argv[2] if len(sys.argv) > 2 else None
    extract_tables(input_file, output_dir)
