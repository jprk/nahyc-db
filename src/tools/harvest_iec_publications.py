"""One-time (idempotent, re-runnable) harvest of every IEC technical
committee's published-standards catalog, via headless Chromium
(Playwright) — doc/PLAN.md §25, 2026-09-17.

`www.iec.ch` sits behind an AWS WAF Bot Control "challenge" action that a
plain HTTP client (`requests`/`curl`) cannot pass — confirmed live before
writing this, and why every other `src/sites/` module's plain-requests
approach doesn't work here. Headless Chromium DOES pass it, but only
once the system's real Chrome runtime libraries are present — this
development environment needed `sudo .venv/bin/playwright install-deps
chromium` as a one-time environment setup step (a real, root-requiring
step; not something this script or the normal pipeline ever repeats).

IEC designations don't encode which of the ~224 technical committees
publishes them, so there is no per-designation live search the way
`normoff.py`/`eiga.py` do it — instead, this harvests EVERY committee's
own "Projects/Publications" export: an XLSX list of Reference/Edition/
Corrigenda-IS/Date/Title/Language/Description, triggered on the site by
a `javascript:openPopup(...)` link that resolves to a plain
`f?p=103:75:0::::FSP_ORG_ID,FSP_LANG_ID,FSP_EXPORT:<org_id>,25,XLSX`
URL once the committee's own `FSP_ORG_ID` is known (read straight off
the committee list table — no popup-clicking needed, Playwright can
navigate to the resolved URL directly and capture the download).

Idempotent (same convention as `fetch_authoritative_metadata.py`'s
cache): each committee's raw parsed rows are written to
`data/iec_committees/<org_id>.json`; a committee already harvested is
skipped unless `--force`. `--build-index-only` skips the harvest
entirely and just re-aggregates `data/iec_publications_index.json` from
whatever committee files already exist — useful after fixing a matching
bug without re-running the (slow, browser-driven) harvest itself.
`src/sites/iec.py` reads only the final index; it never touches
Playwright or the network.

Usage:
    .venv/bin/python src/tools/harvest_iec_publications.py [--force] [--limit N]
    .venv/bin/python src/tools/harvest_iec_publications.py --build-index-only
"""
import argparse
import html
import json
import pathlib
import re
import sys

import openpyxl
from bs4 import BeautifulSoup
from playwright.sync_api import sync_playwright

BASE_DIR = pathlib.Path(__file__).resolve().parent
REPO_ROOT = BASE_DIR.parent.parent

COMMITTEE_LIST_URL = "https://www.iec.ch/technical-committees-and-subcommittees"
USER_AGENT = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36")

COMMITTEES_DIR = REPO_ROOT / "data" / "iec_committees"
INDEX_PATH = REPO_ROOT / "data" / "iec_publications_index.json"

# Reference strings that are an amendment/corrigendum of a base document
# (e.g. "IEC 60050-102:2007/AMD1:2017") never carry their own description
# — confirmed live across all of TC 1's 265 publications, 0 exceptions —
# so the index only ever keeps a BASE reference (no "/AMDn"/"/CORn"
# segment), never one of these.
_AMENDMENT_OR_CORRIGENDUM_RE = re.compile(r"/(AMD|COR)\d")

# "IEC 60050-102:2007" -> "IEC 60050-102" — the edition year (and any
# trailing amendment/corrigendum chain after it) is not part of the
# document's own identity, just this particular row's.
_EDITION_SUFFIX_RE = re.compile(r":\d{4}.*$")


def reference_base(reference):
    return _EDITION_SUFFIX_RE.sub("", (reference or "").strip()).strip()


def is_amendment_or_corrigendum(reference):
    return bool(_AMENDMENT_OR_CORRIGENDUM_RE.search(reference or ""))


def clean_description(text):
    """The XLSX cells carry real HTML markup — not just `<br/>`, but full
    tags and comments (e.g. `<!-- NEW! --><a href="...">IEC 60034-1:2026
    RLV</a>` announcing a newer edition is out) — and a stray
    `_x000D_` (Excel's own escaped-carriage-return artifact, never real
    content). BeautifulSoup's plain-text extraction handles the general
    case correctly (an `<a>` tag's own visible text is kept, comments and
    the tags themselves are dropped) rather than special-casing every
    tag doc/PLAN.md's other extractors happen to hit. Returns None for
    anything that ends up empty, same "empty means no description, not
    an empty string" convention as every other site module."""
    if not text:
        return None
    text = text.replace("_x000D_", " ")
    text = BeautifulSoup(text, "html.parser").get_text(" ")
    text = html.unescape(text)
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


def parse_committee_workbook(path):
    """Every row of a committee's exported XLSX as a plain dict list —
    [] if the file has no data rows at all (a committee with zero
    publications, which does happen for a newly-formed one)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    header = [(c or "").strip() for c in rows[0]]
    out = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        record = dict(zip(header, row))
        out.append({
            "reference": (record.get("Reference") or "").strip(),
            "edition": record.get("Edition"),
            "date": str(record.get("Date")) if record.get("Date") else None,
            "title": (record.get("Title") or "").strip() or None,
            "language": (record.get("Language") or "").strip() or None,
            "description": clean_description(record.get("Description")),
        })
    return out


def fetch_committee_list(page):
    """Every committee's (name, FSP_ORG_ID) from the live table — 224
    confirmed live 2026-09-17. Skips a row with no org-id link at all
    (the second, "disbanded"-history table on the same page has none)."""
    page.goto(COMMITTEE_LIST_URL, timeout=30000, wait_until="networkidle")
    page.wait_for_timeout(2000)
    table = page.query_selector_all("table")[0]
    committees = []
    for row in table.query_selector_all("tr")[1:]:
        if not row.query_selector_all("td"):
            continue
        name = row.query_selector("td").inner_text().strip()
        link = row.query_selector("a[href*='FSP_ORG_ID']")
        if link is None:
            continue
        href = link.get_attribute("href") or ""
        m = re.search(r"FSP_ORG_ID:?,?(\d+)$", href)
        if not m:
            continue
        committees.append({"name": name, "org_id": m.group(1)})
    return committees


def export_url(org_id):
    return (f"https://www.iec.ch/dyn/www/f?p=103:75:0::::"
            f"FSP_ORG_ID,FSP_LANG_ID,FSP_EXPORT:{org_id},25,XLSX")


def download_committee_publications(page, org_id, tmp_path):
    """Downloads one committee's publications XLSX to `tmp_path`. Returns
    True on success, False on any failure (e.g. a committee with an
    export that times out) — never raises, so one bad committee doesn't
    abort the whole harvest.

    `page.goto()` on this URL always raises "Download is starting" —
    Playwright's own signal that navigation turned into a file download
    rather than a page load, not a real failure. The download itself
    still fires and is captured by `expect_download()` regardless;
    swallowed here so it isn't mistaken for one."""
    try:
        with page.expect_download(timeout=20000) as dl_info:
            try:
                page.goto(export_url(org_id), timeout=30000, wait_until="commit")
            except Exception:
                pass
        dl_info.value.save_as(tmp_path)
        return True
    except Exception:
        return False


def build_index():
    """Aggregates every `data/iec_committees/<org_id>.json` into the
    final designation -> description lookup, `data/iec_publications
    _index.json`. Keyed by `reference_base()`, never a full reference
    with its own edition year — a corpus designation doesn't carry an
    edition either. When a base reference appears more than once (a
    superseded edition still listed, or the same publication co-owned by
    more than one committee), the entry with the latest `date` wins —
    same "prefer most recent" rule `normoff.py`'s `resolve()` already
    uses, for the same reason (a superseded edition's own scope can still
    differ from the current one, but the current one is the better
    default). Amendment/corrigendum rows never enter the index at all —
    confirmed live they never carry a description of their own."""
    index = {}
    for committee_file in sorted(COMMITTEES_DIR.glob("*.json")):
        with open(committee_file, "r", encoding="utf-8") as f:
            rows = json.load(f)
        for row in rows:
            reference = row["reference"]
            if not reference or is_amendment_or_corrigendum(reference):
                continue
            if not row.get("description"):
                continue
            base = reference_base(reference)
            existing = index.get(base)
            if existing is None or (row.get("date") or "") > (existing.get("date") or ""):
                index[base] = row
    with open(INDEX_PATH, "w", encoding="utf-8") as f:
        json.dump(index, f, ensure_ascii=False, indent=2, sort_keys=True)
    return index


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--force", action="store_true",
                         help="re-download a committee even if already harvested")
    parser.add_argument("--limit", type=int, default=None,
                         help="harvest at most N committees (for a manual smoke test)")
    parser.add_argument("--build-index-only", action="store_true",
                         help="skip the harvest, just rebuild the index from what's already "
                              "in data/iec_committees/")
    args = parser.parse_args()

    COMMITTEES_DIR.mkdir(parents=True, exist_ok=True)

    if not args.build_index_only:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            context = browser.new_context(user_agent=USER_AGENT, accept_downloads=True)
            page = context.new_page()

            committees = fetch_committee_list(page)
            print(f"{len(committees)} committees found.")

            harvested = 0
            failed = []
            for committee in committees:
                if args.limit is not None and harvested >= args.limit:
                    break
                org_id = committee["org_id"]
                out_path = COMMITTEES_DIR / f"{org_id}.json"
                if out_path.exists() and not args.force:
                    continue
                tmp_path = COMMITTEES_DIR / f"{org_id}.xlsx"
                ok = download_committee_publications(page, org_id, tmp_path)
                if not ok:
                    failed.append(committee["name"])
                    continue
                rows = parse_committee_workbook(tmp_path)
                tmp_path.unlink()
                with open(out_path, "w", encoding="utf-8") as f:
                    json.dump(rows, f, ensure_ascii=False, indent=2)
                harvested += 1
                print(f"[{committee['name']}] {len(rows)} publication(s)")

            browser.close()

        print(f"\nHarvested {harvested} committee(s) this run.")
        if failed:
            print(f"{len(failed)} committee(s) failed (export timed out or no download): "
                  f"{', '.join(failed)}")

    index = build_index()
    with_description = sum(1 for v in index.values() if v.get("description"))
    print(f"Index rebuilt: {len(index)} base reference(s), "
          f"{with_description} with a description -> "
          f"{INDEX_PATH.relative_to(REPO_ROOT)}")


if __name__ == "__main__":
    main()
