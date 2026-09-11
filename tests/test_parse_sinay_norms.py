import pathlib
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent / "src" / "tools"))

from parse_sinay_norms import (
    classify_jurisdikce, _pdf_row_to_record, is_placeholder_designation,
    merge_designation_continuations, parse_xlsx, _XLSX_SHEET, _XLSX_HEADER_ROW,
)


class ClassifyJurisdikceTestCase(unittest.TestCase):
    def test_stn_designation_is_slovak(self):
        self.assertEqual(classify_jurisdikce("STN EN 17124/ - 2022.06"), "SK")

    def test_german_bodies_are_de(self):
        for kategorie in ("DVGW-Arbeitsblatt (DVGW G-TK-1-4)", "Norm (DIN)", "Technischer Bericht (VDI)",
                          "Merkblatt (DASt)", "Technische Regel (BAuA)"):
            self.assertEqual(classify_jurisdikce("G 102-2 (A)", kategorie), "DE", msg=kategorie)

    def test_international_bodies_are_mezinarodni(self):
        self.assertEqual(classify_jurisdikce("22734", "ISO/TC 197 Hydrogen Technologies"), "mezinárodní")
        self.assertEqual(classify_jurisdikce("60092-506", "IEC/TC 18"), "mezinárodní")

    def test_cen_cenelec_and_eiga_are_eu(self):
        self.assertEqual(classify_jurisdikce("13480-9", "CEN/TC 54"), "EU")
        self.assertEqual(classify_jurisdikce("IGC 15", "EIGA"), "EU")

    def test_bare_en_designation_with_no_national_prefix_is_eu(self):
        self.assertEqual(classify_jurisdikce("EN 1717"), "EU")
        self.assertEqual(classify_jurisdikce("prEN 13480-9"), "EU")
        self.assertEqual(classify_jurisdikce("FprEN 62282-3-400"), "EU")

    def test_stn_en_prefix_is_still_slovak_not_eu(self):
        # STN must win over the bare-EN fallback — checked first in the
        # marker list, and this designation does have a national prefix.
        self.assertEqual(classify_jurisdikce("STN EN 17124/ - 2022.06"), "SK")

    def test_din_designation_wins_over_international_committee_in_kategorie(self):
        # Real, previously-mismatched cases (doc/PLAN.md Step 1 follow-up
        # #12): a DIN-adopted standard's own catalog entry often cites the
        # international/European committee that originated it -- that must
        # not override the designation's own DE (German) jurisdiction, the
        # same way STN already wins over a cited international committee.
        self.assertEqual(classify_jurisdikce("DIN EN IEC 60079-11", "Norm (IEC/TC31)"), "DE")
        self.assertEqual(classify_jurisdikce("DIN EN 10216-2", "Norm (CEN/TC 459/SC 10/WG 1)"), "DE")

    def test_stn_still_wins_over_a_cited_german_mirror_committee(self):
        # The reordering that fixed the DIN case above must not disturb
        # STN's own, already-correct precedence when an STN-adopted
        # standard's catalog entry cites the German mirror committee.
        self.assertEqual(
            classify_jurisdikce("STN EN 13096/ – 2004.12", "Norm (DIN-Normenausschuss Druckgasanlagen)"),
            "SK")

    def test_us_bodies(self):
        for kategorie in ("CGA - Compressed Gas Association", "ASTM International"):
            self.assertEqual(classify_jurisdikce("SOME-CODE", kategorie), "US", msg=kategorie)

    def test_unrecognized_us_body_is_neurceno_not_a_guess(self):
        # "US Air Force" isn't in the marker list — a deliberate residual,
        # not every conceivable issuing body is worth chasing (see
        # doc/PLAN.md Step 1 follow-up #8).
        self.assertEqual(classify_jurisdikce("A-A-59874", "US Air Force (Air Force Life Cycle Management Center)"),
                          "neurčeno")

    def test_unrecognized_body_is_neurceno_not_a_guess(self):
        self.assertEqual(classify_jurisdikce("SEP 1970", "SEP – Stahl-Eisen-Prüfblätter"), "neurčeno")
        self.assertEqual(classify_jurisdikce("bez označenia", ""), "neurčeno")


class PdfRowToRecordTestCase(unittest.TestCase):
    def test_adopted_stn_gets_stn_kategorie_and_platnost_from_date(self):
        row = {
            "designation": "STN EN 17124/ - 2022.06",
            "title": "Vodíkové palivo - Špecifikácia výrobku",
            "url": "https://normy.normoff.gov.sk/",
            "note": "Prijatá v sústave STN",
        }
        record = _pdf_row_to_record(row)
        self.assertEqual(record["Kategorie"], "STN (prijatá v sústave)")
        self.assertEqual(record["Platnost"], "od 06/2022")
        self.assertEqual(record["Jurisdikce"], "SK")
        self.assertEqual(record["Link"], "https://normy.normoff.gov.sk/")

    def test_foreign_body_note_becomes_kategorie(self):
        row = {
            "designation": "CGA H-10",
            "title": "Combustion safety for steam reformer",
            "url": "https://www.cganet.com/",
            "note": "CGA - Compressed Gas Association; US",
        }
        record = _pdf_row_to_record(row)
        self.assertEqual(record["Kategorie"], "CGA - Compressed Gas Association; US")
        self.assertEqual(record["Jurisdikce"], "US")

    def test_url_with_embedded_wrap_space_is_stripped(self):
        row = {"designation": "ISO 21011", "title": "Cryogenic vessels",
               "url": "https://www.iso.org/standards.h tml", "note": ""}
        record = _pdf_row_to_record(row)
        self.assertEqual(record["Link"], "https://www.iso.org/standards.html")

    def test_placeholder_designation_becomes_empty_znacka(self):
        # A real case found in the source PDF: the designation column
        # literally holds Slovak text meaning "no designation available"
        # for a guidance document that simply doesn't have a formal
        # standard number -- taking it literally as a real znacka
        # silently blocked a legitimate same-title merge (doc/PLAN.md
        # Step 1 follow-up #10/#11).
        row = {"designation": "Bez označenia", "title": "BVEG Leitfaden Bohrungsintegrität",
               "url": "https://www.bveg.de/", "note": "BVEG"}
        record = _pdf_row_to_record(row)
        self.assertEqual(record["Značka"], "")
        # the record itself must still be kept -- it's real content, just
        # without a formal reference number (see parse_pdf's title-only gate).
        self.assertEqual(record["Název"], "BVEG Leitfaden Bohrungsintegrität")


class IsPlaceholderDesignationTestCase(unittest.TestCase):
    def test_known_placeholders_case_insensitive(self):
        self.assertTrue(is_placeholder_designation("bez označenia"))
        self.assertTrue(is_placeholder_designation("Bez označenia"))
        self.assertTrue(is_placeholder_designation("keine Nummer vorhanden"))
        self.assertTrue(is_placeholder_designation("  keine nummer vorhanden  "))

    def test_real_designation_is_not_a_placeholder(self):
        self.assertFalse(is_placeholder_designation("ISO 14687"))
        self.assertFalse(is_placeholder_designation(""))
        self.assertFalse(is_placeholder_designation(None))


class MergeDesignationContinuationsTestCase(unittest.TestCase):
    def test_date_fragment_continuation_is_merged_with_a_space(self):
        lines = [(100, "STN EN ISO 11114-1/ –"), (112, "2020.12")]
        result = merge_designation_continuations(lines)
        self.assertEqual(result, [(100, "STN EN ISO 11114-1/ – 2020.12")])

    def test_short_numeric_continuation_after_hyphen_gets_no_space(self):
        # Real case: "Sandia Report SAND2012-" / "7321" wrapped onto two
        # lines in the source PDF -- found via the same duplicate-title
        # audit that caught the edition-date-suffix bug.
        lines = [(100, "Sandia Report SAND2012-"), (112, "7321")]
        result = merge_designation_continuations(lines)
        self.assertEqual(result, [(100, "Sandia Report SAND2012-7321")])

    def test_short_paren_continuation_gets_a_space(self):
        lines = [(100, "UL Standard (UL 125, Edition"), (112, "1)")]
        result = merge_designation_continuations(lines)
        self.assertEqual(result, [(100, "UL Standard (UL 125, Edition 1)")])

    def test_unrelated_short_line_is_not_merged(self):
        # A short line that isn't purely digits/closing-punctuation (e.g.
        # a genuinely new, short designation) must start its own row.
        lines = [(100, "Sandia Report SAND2012-"), (112, "ADR")]
        result = merge_designation_continuations(lines)
        self.assertEqual(result, [(100, "Sandia Report SAND2012-"), (112, "ADR")])

    def test_single_line_is_unchanged(self):
        lines = [(100, "ISO 14687")]
        self.assertEqual(merge_designation_continuations(lines), lines)

    def test_empty_input(self):
        self.assertEqual(merge_designation_continuations([]), [])


class ParseXlsxPlaceholderTestCase(unittest.TestCase):
    """Builds a small in-memory workbook matching parse_xlsx's expected
    layout (header row 14, data from row 15, sheet "NRM H2_Bestandsanalyse")
    to verify a row with a placeholder-only designation is kept (not
    dropped) when it has a real title, and correctly gets an empty znacka."""

    def _make_workbook(self, rows):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = _XLSX_SHEET
        # Only the columns parse_xlsx actually reads need real headers;
        # the AG-flag range (24-63) is left blank, which parse_xlsx
        # tolerates (produces an empty klicova_slova list).
        ws.cell(row=_XLSX_HEADER_ROW, column=1, value="Publikationsform")
        for row_offset, row in enumerate(rows):
            r = 15 + row_offset
            ws.cell(row=r, column=1, value=row.get("pub_form", ""))
            ws.cell(row=r, column=2, value=row.get("foreign_designation", ""))
            ws.cell(row=r, column=3, value=row.get("stn_designation", ""))
            ws.cell(row=r, column=5, value=row.get("stn_title", ""))
            ws.cell(row=r, column=9, value=row.get("title_de", ""))
            ws.cell(row=r, column=10, value=row.get("title_en", ""))
        return wb

    def _parse(self, rows):
        wb = self._make_workbook(rows)
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = pathlib.Path(tmp_dir) / "sinay_test.xlsx"
            wb.save(path)
            return parse_xlsx(path=path)

    def test_placeholder_designation_with_real_title_is_kept(self):
        records = self._parse([
            {"foreign_designation": "keine Nummer vorhanden",
             "title_de": "BVEG Leitfaden Bohrungsintegrität"},
        ])
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["Značka"], "")
        self.assertEqual(records[0]["Název"], "BVEG Leitfaden Bohrungsintegrität")

    def test_row_with_no_designation_and_no_title_is_dropped(self):
        self.assertEqual(self._parse([{}]), [])

    def test_real_designation_row_still_works(self):
        records = self._parse([
            {"stn_designation": "STN EN 17124", "stn_title": "Vodíkové palivo"},
        ])
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["Značka"], "STN EN 17124")


if __name__ == "__main__":
    unittest.main()
